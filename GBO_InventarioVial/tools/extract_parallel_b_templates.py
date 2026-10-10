"""Read-only contract extraction; originals are copied byte-for-byte as canonical assets."""
import hashlib, json, shutil, zipfile
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SOURCES = {
    'LOSA': Path('C:/Users/antho/Downloads/FICHA SCAP - GBO modelo - LOSA.xlsx'),
    'MODULAR': Path('C:/Users/antho/Downloads/FICHA SCAP - GBO modelo - MODULAR.xlsx'),
    'SIC17': Path('C:/Users/antho/Downloads/SIC17_MO-104.xlsx'),
}

def extract(kind, source):
    wb = openpyxl.load_workbook(source, data_only=False)
    result = {'source':source.name, 'sha256':hashlib.sha256(source.read_bytes()).hexdigest(),
              'definedNames':{k:v.attr_text for k,v in wb.defined_names.items()}, 'sheets':[]}
    for s in wb:
        cells = [{'cell':c.coordinate,'value':c.value,'type':c.data_type,'style':c.style_id}
                 for row in s for c in row if c.value is not None]
        validations = [{'range':str(d.sqref),'type':d.type,'formula1':d.formula1,'formula2':d.formula2,
                        'allowBlank':d.allowBlank} for d in s.data_validations.dataValidation]
        result['sheets'].append({'name':s.title,'maxRow':s.max_row,'maxColumn':s.max_column,
            'cells':cells,'merges':[str(m) for m in s.merged_cells.ranges],
            'validations':validations,'rowDimensions':{str(k):{'height':v.height,'hidden':v.hidden} for k,v in s.row_dimensions.items()},
            'columnDimensions':{k:{'width':v.width,'hidden':v.hidden} for k,v in s.column_dimensions.items()},
            'headerFooter':str(s.HeaderFooter),'printArea':str(s.print_area),
            'images':[{'width':i.width,'height':i.height,'anchorType':type(i.anchor).__name__,
                       'from':vars(i.anchor._from) if hasattr(i.anchor,'_from') else None,
                       'to':vars(i.anchor.to) if hasattr(i.anchor,'to') else None} for i in s._images],
            'formulas':[c for c in cells if c['type']=='f'],
            'sections':[c for c in cells if isinstance(c['value'],str) and any(c['value'].strip().startswith(x) for x in ['A.','B.','C.','D.','E.','F.','G.'])]})
    with zipfile.ZipFile(source) as z:
        result['mediaParts']=[n for n in z.namelist() if n.startswith('xl/media/')]
    path=ROOT/'docs'/ (f'SCAP_{kind}_FIELD_MAP.json' if kind!='SIC17' else 'SIC_WORKBOOK_HEADER_MAP.json')
    path.write_text(json.dumps(result,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
    target=ROOT/'docs/reference'/source.name
    shutil.copyfile(source,target)
    print(kind, [(s['name'],s['maxRow'],len(s['validations'])) for s in result['sheets']])

if __name__=='__main__':
    baseline=ROOT/'output/parallel-b/baseline/source'
    if not baseline.exists():
        baseline.mkdir(parents=True)
        shutil.copytree(ROOT/'app/src',baseline/'app/src')
        shutil.copytree(ROOT/'app/schemas',baseline/'app/schemas')
        shutil.copyfile(ROOT/'app/build.gradle.kts',baseline/'app/build.gradle.kts')
    for kind,source in SOURCES.items(): extract(kind,source)
