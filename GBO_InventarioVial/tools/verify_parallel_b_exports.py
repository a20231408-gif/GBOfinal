"""Read exported XLSX packages independently, never rewrite a source template."""
from pathlib import Path
import hashlib,json,zipfile
import openpyxl

ROOT=Path(__file__).resolve().parents[1]

def verify():
    result={}
    for kind in ['LOSA','MODULAR']:
        source=ROOT/'docs/reference'/f'FICHA SCAP - GBO modelo - {kind}.xlsx'
        output=ROOT/'tmp/parallel-b'/f'SCAP_{kind}.xlsx'
        original=openpyxl.load_workbook(source,data_only=False)
        generated=openpyxl.load_workbook(output,data_only=False)
        assert original.sheetnames==generated.sheetnames
        for old,new in zip(original,generated):
            assert set(map(str,old.merged_cells.ranges))==set(map(str,new.merged_cells.ranges)),old.title
            assert {k:v.width for k,v in old.column_dimensions.items()}=={k:v.width for k,v in new.column_dimensions.items()},old.title
            assert {k:v.height for k,v in old.row_dimensions.items()}=={k:v.height for k,v in new.row_dimensions.items()},old.title
        with zipfile.ZipFile(source) as a,zipfile.ZipFile(output) as b:
            assert b.testzip() is None
            assert a.read('xl/styles.xml')==b.read('xl/styles.xml')
            for logo in ['xl/media/image1.jpeg','xl/media/image2.png']:
                assert a.read(logo)==b.read(logo)
        result[kind]={'sheets':generated.sheetnames,'sheetCount':len(generated.sheetnames),'stylesMergesDimensionsLogoPreserved':True,
                      'sourceSha256':hashlib.sha256(source.read_bytes()).hexdigest()}
    for file in sorted((ROOT/'app/build/sic-export-verification').glob('SIC-*.xlsx')):
        book=openpyxl.load_workbook(file,data_only=False)
        sheet=book.worksheets[0]
        assert sheet['F1'].value.startswith('SERVICIO DE CONSULTORIA')
        assert file.stem in sheet['A10'].value
        assert sheet['C14'].value=='0005',file.name
        assert sheet['C14'].data_type=='s'
        assert len(sheet._images)==2
        result[file.stem]={'ownTitle':sheet['A10'].value,'corporateHeading':True,'originalLogoCount':len(sheet._images),'prPreservedAsText':True}
    path=ROOT/'docs/PARALLEL_B_EXPORT_VERIFICATION.json'
    path.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print('Verified independent XLSX reads:',', '.join(result))

if __name__=='__main__':verify()
