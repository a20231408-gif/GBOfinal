"""Build the editable user manual from its reviewed Spanish source. Use bundled Python."""
from pathlib import Path
import re, json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs/manual-10i'
FIG=OUT/'figuras';FIG.mkdir(parents=True,exist_ok=True)
FONT=Path('C:/Windows/Fonts/arial.ttf')
def font(size):return ImageFont.truetype(str(FONT),size)
def flow():
    im=Image.new('RGB',(1500,270),'white');d=ImageDraw.Draw(im)
    labels=['Guardar\nborrador','Revisar y\nfinalizar','Pendiente\nSubiendo','Sincronizado\nACK confirmado']
    for i,text in enumerate(labels):
        x=170+i*375;d.multiline_text((x,100),text,font=font(32),fill='black',anchor='mm',align='center',spacing=10)
        if i<3:
            d.line((x+115,100,x+230,100),fill='#285563',width=4)
            d.polygon([(x+230,100),(x+216,92),(x+216,108)],fill='#285563')
    d.text((750,224),'Diagrama del proceso   No es una captura de pantalla',font=font(22),fill='#555555',anchor='mm')
    im.save(FIG/'flujo-registro.png')
def slab():
    im=Image.new('RGB',(1100,750),'white');d=ImageDraw.Draw(im)
    left,top,right,bottom=160,65,920,630
    d.rectangle((left,top,right,bottom),outline='#555555',width=3)
    d.line(((left+right)//2,top,(left+right)//2,bottom),fill='#999999',width=2)
    d.line((left,(top+bottom)//2,right,(top+bottom)//2),fill='#999999',width=2)
    for n,xy in [(1,(185,370)),(2,(565,370)),(3,(185,90)),(4,(565,90))]:d.text(xy,f'Losa {n}',font=font(24),fill='#555555')
    def point(x,y):return (left+x/6*(right-left),bottom-y/10*(bottom-top))
    a=point(.5,2);b=point(1.5,1);d.rectangle((*a,*b),outline='#25824a',width=5)
    d.text(point(.3,2.6),'Zona parcial 1',font=font(22),fill='#185b33')
    d.line((*point(.4,3),*point(2.5,3.3)),fill='#b86b14',width=5)
    d.text(point(.3,3.8),'Línea independiente 2',font=font(22),fill='#8b4c00')
    d.text((left,675),'Origen 0 0',font=font(22),fill='black');d.text((850,675),'X en m',font=font(22),fill='black')
    d.text((45,330),'Y en m',font=font(22),fill='black')
    d.text((550,724),'Esquema explicativo con datos ficticios   No es una captura',font=font(22),fill='#555555',anchor='mm')
    im.save(FIG/'croquis-explicativo.png')
def fallback(filename,title,lines):
    im=Image.new('RGB',(1300,410),'white');d=ImageDraw.Draw(im)
    d.text((55,35),title,font=font(32),fill='black')
    for i,line in enumerate(lines):d.text((55,110+i*66),line,font=font(27),fill='black')
    d.text((55,365),'Diagrama explicativo basado en el formulario   No es una captura',font=font(23),fill='#555555')
    im.save(FIG/filename)
flow();slab()
real_captures=[name for name in ['ahuellamiento.png','losa-localizada.png'] if (FIG/name).exists()]
source=(OUT/'MANUAL_USUARIO_GBO_INVENTARIO_VIAL.md').read_text(encoding='utf-8')
source=source.replace('![Renderizador de exportación con datos ficticios](figuras/croquis-localizado.png)','![Localización de dos daños en una losa](figuras/croquis-explicativo.png)')
source=source.replace('Figura 4. Página generada por el renderizador de exportación de Pavimentos con datos ficticios. Es evidencia del croquis, no una captura de una ficha real.','Figura 4. Diagrama explicativo de dos daños independientes en una losa. No es una captura ni una medición contractual.')
for name,title,lines,caption in [
    ('ahuellamiento.png','Ahuellamiento   Código 15',['Área afectada en m²','Profundidad media medida en mm','Severidad según el catálogo   Sin ancho de fisura'],'Figura 2. Diagrama de los datos del ahuellamiento. No es una captura de pantalla.'),
    ('losa-localizada.png','Daño localizado en pavimento rígido',['Seleccionar tipo y severidad','Punto o dos extremos dentro de la losa','Cantidad contractual en losas'],'Figura 3. Diagrama del procedimiento de localización. No es una captura de pantalla.')]:
    if not (FIG/name).exists():
        fallback('diagrama-'+name,title,lines)
        source=source.replace('figuras/'+name,'figuras/diagrama-'+name)
        source=re.sub(r'Figura '+('2' if name.startswith('ahuell') else '3')+r'\.[^\n]+',caption,source)
    else:
        source=source.replace('figuras/diagrama-'+name,'figuras/'+name)
        number='2' if name.startswith('ahuell') else '3'
        source=re.sub(r'Figura '+number+r'\.[^\n]+','Figura '+number+'. Componente real de la aplicación con datos ficticios en pruebas aisladas. No corresponde a la tablet de campo.',source)
source=source.replace('Figura 2. Editor de daño de la aplicación con datos ficticios en entorno aislado. Si no se incluye la imagen, seguir los nombres de campo y el procedimiento descritos; no se sustituye por una interfaz fabricada.','Figura 2. Editor real de la aplicación con datos ficticios en un entorno de pruebas aislado. Se muestra la zona visible tras desplazar el contenido; no corresponde a la tablet de campo.')
(OUT/'MANUAL_USUARIO_GBO_INVENTARIO_VIAL.md').write_text(source,encoding='utf-8')

doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
# The bundled template contains a blue Title rule. Remove inherited paragraph
# borders before authoring so the cover stays plain black as requested.
for border in list(doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
sec.top_margin=Cm(2.0);sec.bottom_margin=Cm(2.0);sec.left_margin=Cm(2.1);sec.right_margin=Cm(2.1)
sec.header_distance=Cm(.8);sec.footer_distance=Cm(.8);sec.different_first_page_header_footer=True
normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10.5)
normal.paragraph_format.line_spacing=1.12;normal.paragraph_format.space_after=Pt(6)
for name,size in [('Title',31),('Subtitle',14),('Heading 1',16),('Heading 2',12),('Heading 3',11)]:
    s=doc.styles[name];s.font.name='Arial';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
    s.font.underline=False;s.paragraph_format.space_before=Pt(14);s.paragraph_format.space_after=Pt(7)
    if name.startswith('Heading'):s.paragraph_format.keep_with_next=True
doc.styles['Caption'].font.size=Pt(9);doc.styles['Caption'].font.color.rgb=RGBColor(55,55,55)
doc.core_properties.title='Manual de usuario de GBO Inventario Vial'
doc.core_properties.subject='Procedimientos de campo y revisión de la candidata 10I'
doc.core_properties.author='GBO';doc.core_properties.keywords='GBO inventario vial SIC SCAP Pavimentos 10I'
doc.core_properties.comments='Edición basada en el código y en datos ficticios de pruebas.'
header=sec.header.paragraphs[0];header.text='GBO Inventario Vial   |   Manual de usuario 10I';header.style='Caption'
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
def field(p,text):
    r=p.add_run();a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'begin');r._r.append(a)
    t=OxmlElement('w:instrText');t.set(qn('xml:space'),'preserve');t.text=' '+text+' ';r._r.append(t)
    a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'end');r._r.append(a)
footer.add_run('Página ');field(footer,'PAGE');footer.add_run(' de ');field(footer,'NUMPAGES')
doc.add_paragraph('GBO Inventario Vial',style='Title')
doc.add_paragraph('Manual de usuario',style='Subtitle')
doc.add_paragraph('Registro de campo   Revisión   Entrega de evidencias')
doc.add_paragraph('Edición 10I\nCandidata 1.4.3\n13 de octubre de 2026')
doc.add_paragraph('Dirigido a ingenieros de campo y administradores autorizados.')
doc.add_paragraph('Documento editable para revisión del propietario antes de la distribución controlada.')
doc.add_page_break();title=doc.add_paragraph('Índice');title.runs[0].font.size=Pt(18);title.runs[0].bold=True
field(doc.add_paragraph(),'TOC \\o "1-1" \\h \\z \\u')
doc.add_page_break()
lines=source.splitlines();i=0
while i<len(lines):
    line=lines[i].strip();i+=1
    if not line or line.startswith('# ') or line.startswith('Edición de campo'):continue
    if line.startswith('## '):doc.add_heading(line[3:],level=1);continue
    image=re.match(r'!\[(.*?)\]\((.*?)\)',line)
    if image:
        path=OUT/image.group(2)
        if path.exists():
            p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
            with Image.open(path) as figure:width,height=figure.size
            p.add_run().add_picture(str(path),width=Cm(min(15.5,16*width/height)))
            drawing=p._p.xpath('.//wp:docPr')
            if drawing:drawing[0].set('descr',image.group(1))
        continue
    if line.startswith('Figura '):doc.add_paragraph(line,style='Caption');continue
    if line.startswith('|'):
        rows=[line]
        while i<len(lines) and lines[i].strip().startswith('|'):rows.append(lines[i].strip());i+=1
        parsed=[[c.strip() for c in row.strip('|').split('|')] for row in rows if not re.fullmatch(r'[| :\-]+',row)]
        table=doc.add_table(rows=0, cols=len(parsed[0]));table.autofit=False
        table.columns[0].width=Cm(4.3);table.columns[1].width=Cm(12.5)
        for ri,row in enumerate(parsed):
            cells=table.add_row().cells
            for ci,text in enumerate(row):
                cells[ci].width=Cm(4.3 if ci==0 else 12.5)
                cells[ci].text=text
                tcPr=cells[ci]._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
                for edge in ['top','bottom','left','right']:
                    el=OxmlElement('w:'+edge);el.set(qn('w:w'),'100');el.set(qn('w:type'),'dxa');m.append(el)
                tcPr.append(m)
                for p in cells[ci].paragraphs:
                    p.paragraph_format.space_after=Pt(3)
                    if ri==0:p.paragraph_format.keep_with_next=True
                    for r in p.runs:r.font.size=Pt(10);r.bold=ri==0
                if ri==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EAEAEA');tcPr.append(shade)
            trPr=table.rows[-1]._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
            if ri==0:trPr.append(OxmlElement('w:tblHeader'))
        doc.add_paragraph();continue
    p=doc.add_paragraph(line)
    if re.match(r'^\d+\. ',line):p.paragraph_format.left_indent=Cm(.45);p.paragraph_format.first_line_indent=Cm(-.45)
settings=doc.settings.element;update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
path=OUT/'MANUAL_USUARIO_GBO_INVENTARIO_VIAL.docx';doc.save(path)
(OUT/'proveniencia-figuras.json').write_text(json.dumps({'actualIsolatedComposableCaptures':real_captures,'diagrams':['flujo-registro.png','croquis-explicativo.png']+['diagrama-'+name for name in ['ahuellamiento.png','losa-localizada.png'] if name not in real_captures],'realTabletScreenshots':False,'realPersonalData':False},indent=2),encoding='utf-8')
print(path)
