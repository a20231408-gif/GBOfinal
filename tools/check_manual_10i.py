"""Scrub generated Office metadata, render every PDF page and record structural QA."""
from pathlib import Path
import json,re,zipfile,io
import xml.etree.ElementTree as ET
from docx import Document
from pypdf import PdfReader
import pypdfium2 as pdfium
import pdfplumber
from PIL import Image,ImageDraw

ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'docs/manual-10i'
docx=OUT/'MANUAL_USUARIO_GBO_INVENTARIO_VIAL.docx';pdf=OUT/'MANUAL_USUARIO_GBO_INVENTARIO_VIAL.pdf'
with zipfile.ZipFile(docx) as z:parts={n:z.read(n) for n in z.namelist()}
core=ET.fromstring(parts['docProps/core.xml'])
for node in core:
    if node.tag.rsplit('}',1)[-1] in ['creator','lastModifiedBy']:node.text='GBO'
parts['docProps/core.xml']=ET.tostring(core,encoding='utf-8',xml_declaration=True)
app=ET.fromstring(parts['docProps/app.xml'])
for node in app:
    if node.tag.rsplit('}',1)[-1] in ['Company','Manager']:node.text='GBO' if node.tag.endswith('Company') else ''
parts['docProps/app.xml']=ET.tostring(app,encoding='utf-8',xml_declaration=True)
buf=io.BytesIO()
with zipfile.ZipFile(buf,'w',zipfile.ZIP_DEFLATED) as z:
    for name,data in parts.items():z.writestr(name,data)
docx.write_bytes(buf.getvalue())
document=Document(docx)
heads=[p.text for p in document.paragraphs if p.style.name=='Heading 1']
assert len(heads)==30,len(heads)
for p in document.paragraphs:
    if p.style.name.startswith('Heading') or p.style.name in ['Title','Subtitle']:
        assert p.style.font.color.rgb is None or str(p.style.font.color.rgb)=='000000'
reader=PdfReader(pdf);assert len(reader.pages)>=12
text='\n'.join(p.extract_text() or '' for p in reader.pages)
for word in ['Ahuellamiento','SCAP','OBSERVACIONES','Sincronizado','Administración','PINTADO OK 57 E2','30 Referencias']:assert word in text,word
assert 'Actualice' not in text and 'Error! Bookmark' not in text and 'Error! Reference' not in text
qa=ROOT/'app/build/cierre-10i/manual-qa';qa.mkdir(parents=True,exist_ok=True)
render=pdfium.PdfDocument(pdf)
for index in range(len(render)):
    page=render[index];bitmap=page.render(scale=1.5)
    bitmap.to_pil().save(qa/f'page-{index+1}.png');bitmap.close();page.close()
render.close()
thumbs=[]
for index in range(len(reader.pages)):
    image=Image.open(qa/f'page-{index+1}.png').convert('RGB');image.thumbnail((250,355))
    tile=Image.new('RGB',(270,390),'#dddddd');tile.paste(image,((270-image.width)//2,10))
    ImageDraw.Draw(tile).text((12,370),str(index+1),fill='black');thumbs.append(tile)
columns=3;rows=(len(thumbs)+columns-1)//columns
sheet=Image.new('RGB',(columns*270,rows*390),'white')
for i,image in enumerate(thumbs):sheet.paste(image,((i%columns)*270,(i//columns)*390))
sheet.save(qa/'contact.png')
outside=[]
with pdfplumber.open(pdf) as p:
    for i,page in enumerate(p.pages):
        for char in page.chars:
            if char['x0']<0 or char['x1']>page.width+.5 or char['top']<0 or char['bottom']>page.height+.5:outside.append(i+1);break
assert not outside,outside
result={'pages':len(reader.pages),'chapters':len(heads),'embeddedImages':len(document.inline_shapes),'renderedAllPages':True,
        'pageImagesReviewed':False,'textOutsidePage':outside,'tocUpdatedBy':'Microsoft Word','renderer':'Microsoft Word PDF export and PDFium page rasterization',
        'canonicalRendererUnavailable':'LibreOffice is not installed on this Windows computer','personalOfficeAuthorScrubbed':True,'source':'Markdown and build_manual_10i.py'}
(ROOT/'docs/campo-10i/evidencia/manual-qa.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
print(json.dumps(result))
